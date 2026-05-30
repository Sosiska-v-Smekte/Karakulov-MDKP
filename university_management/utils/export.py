import openpyxl
from openpyxl.styles import Font, Alignment
from core.models import Payment, WorkTeam, WorkType


def export_payments_to_excel(db, filepath):
    """
    Выгружает всю историю выплат в Excel-файл.
    Идеально для формирования квартальных отчетов бухгалтерии.
    """
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "История выплат"

        # Создаем заголовки
        headers = ["ID Выплаты", "ФИО Сотрудника", "Должность", "Выполненная работа", "Дата выплаты", "Сумма (руб.)"]
        ws.append(headers)

        # Стилизуем заголовки (жирный шрифт, выравнивание)
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        # Запрашиваем данные из БД (только успешные выплаты)
        payments = db.query(Payment).filter(Payment.payment_status == 'Выплачено').all()

        total_amount = 0
        for pay in payments:
            emp = pay.work_team.employee
            emp_name = f"{emp.last_name} {emp.first_name}"
            role = "Преподаватель" if emp.role == 'teacher' else "УВП"
            work_title = pay.work_team.work.title

            ws.append([
                pay.id,
                emp_name,
                role,
                work_title,
                pay.payment_date.strftime("%d.%m.%Y") if pay.payment_date else "",
                float(pay.amount)
            ])
            total_amount += float(pay.amount)

        # Итоговая строка
        ws.append([])
        ws.append(["", "", "", "", "ИТОГО:", total_amount])

        # Автоширина колонок
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[col_letter].width = adjusted_width

        wb.save(filepath)
        return True, "Успешно экспортировано!"
    except Exception as e:
        return False, str(e)


def import_work_types_from_excel(db, filepath):
    """
    Загружает справочник работ и ставок из Excel в базу данных.
    Формат файла (колонки): Название работы | Базовая ставка
    """
    try:
        wb = openpyxl.load_workbook(filepath)
        ws = wb.active

        added_count = 0
        # Пропускаем первую строку (заголовки) и читаем данные
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:  # Если название пустое, пропускаем
                continue

            name = str(row[0]).strip()
            try:
                base_cost = float(row[1])
            except (ValueError, TypeError):
                base_cost = 0.0

            # Проверяем, есть ли уже такая работа, чтобы не дублировать
            existing = db.query(WorkType).filter(WorkType.name == name).first()
            if not existing:
                new_type = WorkType(name=name, base_cost=base_cost, unit="шт/час", is_active=True)
                db.add(new_type)
                added_count += 1

        db.commit()
        return True, f"Успешно импортировано {added_count} новых типов работ."
    except Exception as e:
        db.rollback()
        return False, str(e)